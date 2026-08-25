#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
audit_term_collision.py — 각주(†)로 정의한 용어가 문서 다른 곳에서 다른 뜻으로 또 쓰이는지 찾는다.

왜 있나 (teacher_violence_qol, 2026-08-11)
  지도교수가 포스터를 보고 헷갈려했다: Figure 2 캡션 한 문장 안에서 "combined"가
  **wave 통합**("The combined estimate averages away a change")과 **결과변수 통합**
  ("Combined† = the three items below it") 두 뜻으로 쓰이고 있었다. 정의 자체는 어디서도
  틀리지 않았다 — 정의 안 된 자리에서 같은 단어가 재사용된 게 사고였다. 사후에 grep
  한 번으로 5분 만에 찾았는데, 사전에 이 grep을 절차로 안 돌려서 지적을 받았다.

무엇을 하나
  1. 문서 전체에서 "단어†" 패턴(각주 기호가 바로 붙은 단어)을 찾아 **정의된 용어**로 취급한다.
  2. 그 용어가 각주와 함께 등장하는 모든 블록의 어휘를 모아 "정의 맥락 어휘"로 삼는다.
  3. 같은 용어가 각주 없이 등장하는 모든 곳을 찾아, 그 블록의 어휘가 정의 맥락 어휘와
     하나도 안 겹치면 **[의심]**으로 표시한다 — "이 단어를 정의와 무관하게 일반 영어 단어로
     썼을 가능성"이 높다는 뜻이다. 겹치면 [정상]으로, 참고용으로만 보여준다.

  완전 자동 판정이 아니라 **"사람이 5분 걸려 훑을 걸 5초짜리 목록으로 만드는" 도구**다
  (output_surface.md의 "기계적 훑기, 그다음에 눈으로 한 번 더" 원칙과 같다). [의심]이 0건이면
  통과, 1건이라도 있으면 그 자리만 사람이 읽어서 판단한다.

  각주 기호가 없는 문서(또는 정의를 다른 방식으로 표시하는 문서)를 위해 --term으로 용어를
  직접 지정할 수도 있다. 이 경우 그 용어가 처음 등장하는 블록을 정의 맥락으로 쓴다.

지원 형식: .pptx (슬라이드+표 셀) · .docx (문단+표 셀). .hwpx/.md는 아직 없음.

사용법
  python audit_term_collision.py FILE.pptx
  python audit_term_collision.py FILE.docx --term "Combined" --term "Pooled"
  python audit_term_collision.py FILE.pptx --min-overlap 1   # 기본 0(어휘 하나라도 겹치면 정상)

종료 코드: [의심] 항목이 하나라도 있으면 1, 없으면 0. 정의된 용어가 하나도 없으면 0(정보 없음).
"""
import argparse
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DAGGER = "†"
STOPWORDS = {
    "the", "a", "an", "is", "are", "was", "were", "be", "been", "being",
    "of", "or", "and", "in", "on", "at", "to", "for", "with", "that", "this",
    "these", "those", "as", "by", "from", "it", "its", "not", "no", "but",
    "than", "then", "so", "which", "who", "whom", "each", "any", "all",
    "over", "per", "vs", "same", "other", "both", "one", "two", "three",
}


def _words(text):
    return {w for w in re.findall(r"[A-Za-z]{3,}", text.lower()) if w not in STOPWORDS}


def blocks_pptx(path):
    """[(위치, 텍스트), ...] -- 슬라이드별 텍스트프레임 문단 + 표 셀."""
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for si, slide in enumerate(prs.slides, 1):
        for shi, sh in enumerate(slide.shapes, 1):
            if sh.has_text_frame:
                for pi, para in enumerate(sh.text_frame.paragraphs, 1):
                    t = "".join(r.text for r in para.runs)
                    if t.strip():
                        out.append(("slide %d, shape %d, para %d" % (si, shi, pi), t))
            if getattr(sh, "has_table", False):
                for ri, row in enumerate(sh.table.rows, 1):
                    for ci, cell in enumerate(row.cells, 1):
                        t = cell.text
                        if t.strip():
                            out.append(("slide %d, table row %d col %d" % (si, ri, ci), t))
    return out


def blocks_docx(path):
    from docx import Document
    doc = Document(path)
    out = []
    for pi, p in enumerate(doc.paragraphs, 1):
        if p.text.strip():
            out.append(("paragraph %d" % pi, p.text))
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                t = "\n".join(cp.text for cp in cell.paragraphs)
                if t.strip():
                    out.append(("table %d row %d col %d" % (ti, ri, ci), t))
    return out


def load_blocks(path):
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pptx":
        return blocks_pptx(path)
    if ext == "docx":
        return blocks_docx(path)
    raise SystemExit("지원하지 않는 형식: .%s (.pptx/.docx만 됨)" % ext)


def find_guarded_terms(blocks):
    """'단어†' 패턴에서 정의된 용어를 자동으로 찾는다. 단일 토큰만 -- 여러 단어짜리 용어는
    --term으로 직접 지정해야 한다. (이전 버전은 공백·마침표를 넘어 문장 전체를 삼키는 버그가
    있었다: "same covariates. Combined†"가 "same covariates. Combined" 통째로 잡혔다.)"""
    found = {}
    pat = re.compile(r"\b([A-Za-z]+)\s*" + DAGGER)
    for loc, text in blocks:
        for m in pat.finditer(text):
            term = m.group(1).strip()
            found.setdefault(term.lower(), term)
    return list(found.values())


_SENT_SPLIT = re.compile(r"(?<=[.!?])\s+")


def _sentences(loc, text):
    """블록을 문장 단위로 쪼갠다. (위치, 문장) 목록 -- 문장 경계가 곧 판단 단위다.

    블록(문단·표 셀) 하나에 각주 붙은 등장과 안 붙은 등장이 **같이** 들어 있는 게 바로
    실제 사고 패턴이었다("The combined estimate...Combined† = the three items..."는 한
    캡션, 즉 한 블록이었다). 블록 전체를 하나의 어휘 뭉치로 다루면 그 블록은 항상 자기
    자신과 100% 겹쳐 오탐을 놓친다. 문장으로 쪼개야 "combined estimate" 문장과
    "Combined† = ..." 문장이 서로 다른 어휘 집합으로 분리된다."""
    return [(loc, s) for s in _SENT_SPLIT.split(text) if s.strip()]


def occurrences(blocks, term):
    """term이 등장하는 모든 (위치, 문장, 각주붙었는가) 반환. 대소문자 무시, 단어경계.
    판단 단위는 문장이다 -- 문단 전체가 아니다(위 _sentences 참조)."""
    pat = re.compile(r"\b" + re.escape(term).replace(r"\ ", r"\s+") + r"\b", re.I)
    out = []
    for loc, text in blocks:
        for sloc, sent in _sentences(loc, text):
            for m in pat.finditer(sent):
                end = m.end()
                has_dagger = sent[end:end + 2].lstrip().startswith(DAGGER)
                out.append((sloc, sent, has_dagger))
    return out


def audit_term(blocks, term, min_overlap):
    occs = occurrences(blocks, term)
    if not occs:
        return None

    # 용어 자신의 단어(들)는 정의 어휘·비교 어휘 양쪽에서 뺀다. 안 빼면 "combined"라는
    # 단어 자체가 항상 겹쳐서 모든 등장이 무조건 [정상]으로 나온다 -- 비교가 무의미해진다.
    term_words = _words(term)

    def_keywords = set()
    for loc, text, has_dagger in occs:
        if has_dagger:
            def_keywords |= (_words(text) - term_words)
    if not def_keywords:
        # 각주가 붙은 등장이 하나도 없다 -- --term으로 수동 지정된 경우. 첫 등장을 정의로 쓴다.
        def_keywords = _words(occs[0][1]) - term_words

    # 각 등장(occurrence)마다 처리한다 -- 문장 단위로 쪼갠 뒤라도, 같은 문장 안에 각주
    # 붙은 등장과 안 붙은 등장이 같이 있을 수 있으니 (위치, 텍스트)로 dedup하지 않는다.
    rows = []
    suspect = 0
    for loc, text, has_dagger in occs:
        if has_dagger:
            rows.append(("[정의]", loc, text))
            continue
        overlap = len((_words(text) - term_words) & def_keywords)
        if overlap <= min_overlap:
            rows.append(("[의심]", loc, text))
            suspect += 1
        else:
            rows.append(("[정상]", loc, text))
    return term, def_keywords, rows, suspect


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--term", action="append", default=[],
                    help="각주(†) 없이도 검사할 용어를 직접 지정. 여러 번 지정 가능.")
    ap.add_argument("--min-overlap", type=int, default=0,
                    help="이 값 이하로 겹치면 [의심] (기본 0 = 하나도 안 겹치면 의심)")
    a = ap.parse_args()

    blocks = load_blocks(a.file)
    terms = find_guarded_terms(blocks)
    for t in a.term:
        if t.lower() not in [x.lower() for x in terms]:
            terms.append(t)

    if not terms:
        print("%s -- 각주(†)로 정의된 용어를 찾지 못했다. --term으로 직접 지정하지 않으면 "
              "검사할 게 없다." % a.file)
        return 0

    total_suspect = 0
    for term in terms:
        res = audit_term(blocks, term, a.min_overlap)
        if res is None:
            continue
        term, def_keywords, rows, suspect = res
        total_suspect += suspect
        print("\n=== '%s' (%d회 등장, 정의 맥락 어휘: %s) ==="
              % (term, len(rows), ", ".join(sorted(def_keywords)) or "(없음)"))
        for mark, loc, text in rows:
            snippet = text.strip().replace("\n", " ")
            if len(snippet) > 100:
                snippet = snippet[:100] + "..."
            print("  %-6s %-28s %s" % (mark, loc, snippet))

    print("\n" + "-" * 60)
    if total_suspect:
        print("[의심] %d건 -- 정의와 다른 뜻으로 쓰였을 가능성. 눈으로 확인할 것." % total_suspect)
    else:
        print("의심 사례 없음.")
    return 1 if total_suspect else 0


if __name__ == "__main__":
    sys.exit(main())
