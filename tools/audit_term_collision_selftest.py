#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""audit_term_collision.py의 정오탐을 진짜 사고 사례로 증명한다.

teacher_violence_qol 포스터에서 실제로 있었던 문장(수정 전/후)을 그대로 재현해, 이 도구가
그 사고를 잡아냈을지와 정상 문서를 오탐하지 않는지 둘 다 확인한다. 실제 .pptx/.docx 프로젝트
없이 합성 파일로 테스트한다(poster_kit_selftest.py와 같은 패턴).

Run: python audit_term_collision_selftest.py
"""
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_term_collision as atc


def make_pptx(caption_text, other_blocks=()):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(9), Inches(2))
    box.text_frame.text = caption_text
    for t in other_blocks:
        b = slide.shapes.add_textbox(Inches(0), Inches(2), Inches(9), Inches(1))
        b.text_frame.text = t
    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    tmp.close()
    prs.save(tmp.name)
    return tmp.name


def test_catches_real_bug():
    """수정 전 실제 문장 -- 'combined'가 wave 통합과 결과변수 통합 두 뜻으로 쓰였다."""
    bad_caption = (
        "Figure 2. The combined estimate averages away a change. School teachers are "
        "indistinguishable from other white-collar workers until 2023. Combined† = "
        "the three items below it. Waves with fewer than three exposed teachers are not estimated."
    )
    path = make_pptx(bad_caption)
    try:
        blocks = atc.load_blocks(path)
        terms = atc.find_guarded_terms(blocks)
        assert "Combined" in terms, "Combined† 패턴을 못 찾음: %r" % terms
        term, def_kw, rows, suspect = atc.audit_term(blocks, "Combined", min_overlap=0)
        assert suspect >= 1, "실제 사고를 [의심]으로 못 잡음 -- rows=%r" % (rows,)
        marks = {loc: mark for mark, loc, text in rows}
        print("PASS: test_catches_real_bug (suspect=%d)" % suspect)
    finally:
        os.unlink(path)


def test_clean_after_fix():
    """수정 후 실제 문장 -- wave 통합은 'pooled', 결과변수는 'Combined'로 분리됐다. 의심 0건이어야."""
    good_caption = (
        "Figure 2. The pooled estimate averages away a change. School teachers are "
        "indistinguishable from other white-collar workers until 2023. Combined† = "
        "the three items below it. Waves with fewer than three exposed teachers are not estimated."
    )
    other = "† Combined = verbal abuse, threats or unwanted sexual attention, the three items."
    path = make_pptx(good_caption, other_blocks=[other])
    try:
        blocks = atc.load_blocks(path)
        term, def_kw, rows, suspect = atc.audit_term(blocks, "Combined", min_overlap=0)
        assert suspect == 0, "정상 문서인데 오탐: rows=%r" % (rows,)
        print("PASS: test_clean_after_fix")
    finally:
        os.unlink(path)


def test_legitimate_bare_reference_not_flagged():
    """각주 없이 같은 뜻으로 재언급하는 정상적인 문장 -- 정의 어휘를 최소 하나는 공유한다면
    (실제 이 프로젝트의 재언급 문장들이 다 그렇다: "Combined is the union of..." 식으로 정의
    단어를 되풀이한다) 의심 처리하면 안 된다."""
    caption = (
        "Combined† is the union of verbal abuse, threats and unwanted sexual attention. "
        "Combined includes threats reported by teachers in the past month."
    )
    path = make_pptx(caption)
    try:
        blocks = atc.load_blocks(path)
        term, def_kw, rows, suspect = atc.audit_term(blocks, "Combined", min_overlap=0)
        assert suspect == 0, "정상적인 무각주 재언급을 오탐: rows=%r" % (rows,)
        print("PASS: test_legitimate_bare_reference_not_flagged")
    finally:
        os.unlink(path)


def test_known_limitation_pure_paraphrase_overflags():
    """알려진 한계를 문서화하는 테스트다 -- 실패를 기대하는 게 아니라, 정의 어휘를 하나도
    안 쓰는 완전한 바꿔말하기는 이 도구가 [의심]으로 과잉 표시한다는 걸 명시적으로 남겨둔다.
    완전한 의미이해 없이는 못 피하는 한계이고, "놓치는 것보다 과잉 표시가 낫다"는 설계
    원칙상 의도된 동작이다 -- 사람이 그 한 줄을 보고 2초 만에 넘기면 그만이다."""
    caption = (
        "Combined† is the union of verbal abuse, threats and unwanted sexual attention. "
        "Combined captures any of the three items in the past month."
    )
    path = make_pptx(caption)
    try:
        blocks = atc.load_blocks(path)
        term, def_kw, rows, suspect = atc.audit_term(blocks, "Combined", min_overlap=0)
        assert suspect == 1, ("이 한계 테스트 자체가 깨졌다 -- 동작이 바뀌었으면 문서화도 "
                              "같이 갱신할 것: rows=%r" % (rows,))
        print("PASS: test_known_limitation_pure_paraphrase_overflags (의도된 과잉표시 확인)")
    finally:
        os.unlink(path)


def test_no_guarded_terms_is_clean_exit():
    """각주 자체가 없는 문서는 검사할 게 없다는 안내만 하고 exit 0."""
    path = make_pptx("A plain caption with no daggers at all.")
    try:
        blocks = atc.load_blocks(path)
        terms = atc.find_guarded_terms(blocks)
        assert terms == [], "각주 없는데 용어를 찾음: %r" % terms
        print("PASS: test_no_guarded_terms_is_clean_exit")
    finally:
        os.unlink(path)


def test_no_sentence_boundary_crossing():
    """이전 버전의 실제 버그: 마침표를 건너뛰어 문장 전체를 용어로 삼켰다."""
    caption = "Adjusted for the same covariates. Combined† = the three items below it."
    path = make_pptx(caption)
    try:
        blocks = atc.load_blocks(path)
        terms = atc.find_guarded_terms(blocks)
        assert terms == ["Combined"], "문장 경계를 건너뛰어 잘못된 용어를 잡음: %r" % terms
        print("PASS: test_no_sentence_boundary_crossing")
    finally:
        os.unlink(path)


if __name__ == "__main__":
    test_catches_real_bug()
    test_clean_after_fix()
    test_legitimate_bare_reference_not_flagged()
    test_known_limitation_pure_paraphrase_overflags()
    test_no_guarded_terms_is_clean_exit()
    test_no_sentence_boundary_crossing()
    print("\nALL PASS")
