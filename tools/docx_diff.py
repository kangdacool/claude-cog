#!/usr/bin/env python3
"""Compare two .docx across EVERY axis that can carry an edit.

Why this exists (2026-08-19): a delivered brief came back hand-edited.
I diffed `doc.paragraphs`, saw 41/41 identical, and reported "prose unchanged -
only the font size differs". Wrong: `doc.paragraphs` does NOT include table cell
text, and the user had also restructured a table (long -> wide pivot with a
merged span header). The user caught it, not me.

"The prose did not change" is not "the content did not change."

Axes compared here:
  1. prose        - doc.paragraphs
  2. tables       - cell text, row/col counts, merge shape
  3. geometry     - column widths, image sizes, section margins
  4. type sizes   - measured from a rendered PDF, not from XML

Axis 4 needs rendering because `w:sz` is frequently ABSENT on runs (size is
inherited from the style), so reading XML reports None and hides a real change.
Only the rendered PDF tells you what the reader actually sees.

Usage:
    python docx_diff.py OLD.docx NEW.docx            # axes 1-3
    python docx_diff.py OLD.docx NEW.docx --render   # + axis 4 (needs Word COM)

Exit code 1 if any axis differs, else 0.
"""
import argparse
import difflib
import os
import re
import sys
import tempfile
from collections import Counter

from docx import Document


# ── helpers ───────────────────────────────────────────────────────────────────
def emu_cm(v):
    return None if v is None else round(v / 360000, 2)


def prose(path):
    """Paragraph text, INCLUDING runs inside <w:ins> (tracked insertions).

    ⛔ 2026-09-11: python-docx's `paragraph.text` reads only direct <w:r> children, so a
    document with tracked changes reports every inserted paragraph as an empty string —
    axis 1 was silently under-reporting on exactly the documents most likely to be diffed.
    Deleted runs (<w:del>/<w:delText>) stay OUT: the diff should show the document as it
    reads with changes applied.
    """
    from docx.oxml.ns import qn
    out = []
    for p in Document(path).paragraphs:
        txt = "".join(n.text or "" for n in p._p.iter(qn("w:t")))
        if txt.strip():
            out.append(txt.strip())
    return out


def tables(path):
    """Cell text plus dimensions. Merged cells repeat their text, which is what
    makes a merge visible here at all."""
    out = []
    for ti, t in enumerate(Document(path).tables):
        out.append(f"### TABLE {ti}  rows={len(t.rows)} cols={len(t.columns)}")
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace("\n", "\\n").strip() for c in row.cells]
            out.append(f"T{ti} r{ri}: " + " | ".join(cells))
    return out


def geometry(path):
    """Keyed, so it can be compared item by item.

    NOT a unified diff: with 15+ tables the context lines around a hunk look just
    like unchanged items, and a changed table outside the context window is simply
    absent. Both misled a reader into "only T5 changed" when two tables had moved.
    Every key is compared explicitly and only real changes print.
    """
    d = Document(path)
    out = {}
    s = d.sections[0]
    out["page"] = (f"{emu_cm(s.page_width)}x{emu_cm(s.page_height)} "
                   f"margins L{emu_cm(s.left_margin)} R{emu_cm(s.right_margin)} "
                   f"T{emu_cm(s.top_margin)} B{emu_cm(s.bottom_margin)}")
    st = d.styles["Normal"]
    out["Normal style"] = f"{st.font.name} {st.font.size.pt if st.font.size else None}pt"
    # Key on index only. Putting the header text in the key made a renamed column
    # look like one item removed plus another added, instead of one value changed.
    # Header text belongs to the tables axis; this axis is geometry alone.
    for ti, t in enumerate(d.tables):
        out[f"T{ti} colwidths"] = str([emu_cm(c.width) for c in t.rows[0].cells])
    for i, sh in enumerate(d.inline_shapes):
        out[f"img{i} size"] = f"{emu_cm(sh.width)}x{emu_cm(sh.height)} cm"
    return out


def render_pdf(docx_path, out_pdf):
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        doc.SaveAs(os.path.abspath(out_pdf), FileFormat=17)
        doc.Close(False)
    finally:
        word.Quit()


def type_profile(pdf_path):
    """Chars rendered at each point size - the only trustworthy size measurement."""
    import fitz
    d = fitz.open(pdf_path)
    c = Counter()
    for page in d:
        for blk in page.get_text("dict")["blocks"]:
            for line in blk.get("lines", []):
                for sp in line["spans"]:
                    if sp["text"].strip():
                        c[round(sp["size"], 1)] += len(sp["text"].strip())
    pages = d.page_count
    d.close()
    return c, pages


def report_seq(name, a, b):
    """Line sequences (prose, table cells) -> unified diff. Returns changed count."""
    if a == b:
        print(f"  [SAME] {name}")
        return 0
    changed = 0
    lines = []
    for line in difflib.unified_diff(a, b, "OLD", "NEW", n=1, lineterm=""):
        if line.startswith(("---", "+++", "@@")):
            continue
        if line.startswith(("+", "-")):
            changed += 1
        lines.append("      " + line[:170])
    print(f"  [DIFF] {name} - {changed} changed line(s)")
    for l in lines:
        print(l)
    return changed


def report_map(name, a, b):
    """Keyed items -> explicit per-key comparison. Returns changed count."""
    keys = list(a) + [k for k in b if k not in a]
    diffs = [(k, a.get(k), b.get(k)) for k in keys if a.get(k) != b.get(k)]
    if not diffs:
        print(f"  [SAME] {name} ({len(keys)} items)")
        return 0
    print(f"  [DIFF] {name} - {len(diffs)} of {len(keys)} item(s) changed")
    for k, va, vb in diffs:
        print(f"      {k}")
        print(f"          OLD {va}")
        print(f"          NEW {vb}")
    return len(diffs)


# ── main ──────────────────────────────────────────────────────────────────────

# ---------------------------------------------------------------------------
# axis 5 — graft fit: does the ADDED prose match the host document's habits?
#
# WHY (2026-09-11, 한 위원회 문서). Six rounds of correction on two inserted
# paragraphs. Every number was right from the start; every defect was at the JOINT.
# The four existing axes ask "what changed". None asks "does what I added belong here".
# Nor do the single-document prose auditors (academic_prose, register, text_consistency):
# they read a document alone, and these defects only exist relative to the host.
#
# WHAT IT CAUGHT, had it existed:
#   · a paragraph opened with 「다만」 — the host uses 다만 three times, ALL paragraph-final
#   · an inserted paragraph at 546 chars where the host section runs 275–459
#
# WHAT IT DELIBERATELY DOES NOT CHECK: whether the added prose «화제»가 겉도는가.
# Tried lexical overlap with the neighbouring paragraphs (and with their last sentences)
# — it fired on the GOOD draft and stayed silent on the bad one. 겉도는 것은 낱말이 아니라
# 화제여서 그렇다. 그건 사람이 본다(korean_prose_grafting.md §4·§5).
#
# NOT a verdict. Korean prose fit is a judgement; this narrows where to look.
# 세부 규율은 feedback/korean_prose_grafting.md.
# ---------------------------------------------------------------------------
CONNECTIVES = ["다만", "그러나", "한편", "반면", "따라서", "그리고", "또한", "즉", "그런데"]


def _sentences(par):
    return [s.strip() for s in re.split(r"(?<=다[.])\s*", par) if s.strip()]


def _conn_positions(host):
    """host 가 각 접속어를 «어디에» 쓰는가 — {접속어: [위치…]}."""
    pos = {}
    for par in host:
        ss = _sentences(par)
        for i, s in enumerate(ss):
            for c in CONNECTIVES:
                if s.startswith(c):
                    where = "문단첫" if i == 0 else ("문단끝" if i == len(ss) - 1 else "중간")
                    pos.setdefault(c, []).append(where)
    return pos


def graft_fit(old_paras, new_paras, window=6):
    """Return (issue_count, lines). Compares each ADDED paragraph against its neighbours."""
    # ⚠ difflib 의 'replace' 는 «고친 문단»도 포함한다. 5자를 끼우거나 한 문장을 바꾼
    #   문단까지 「새로 넣은 글」로 세면 분량·어휘 검사가 통째로 오탐이 된다
    #   (2026-09-11 실측: 4건 중 2건이 그것이었다). 원본과 많이 닮았으면 «편집»으로 본다.
    sm = difflib.SequenceMatcher(None, old_paras, new_paras, autojunk=False)
    added = []            # (new_index, text, host_before, host_after)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag not in ("insert", "replace"):
            continue
        for j in range(j1, j2):
            if tag == "replace":
                kin = max((difflib.SequenceMatcher(None, old_paras[i], new_paras[j]).ratio()
                           for i in range(i1, i2)), default=0.0)
                if kin > 0.5:
                    continue          # 고친 문단 — 이 축의 대상이 아니다
            added.append((j, new_paras[j], old_paras[max(0, i1 - 1):i1],
                          old_paras[i2:i2 + 1]))
    if not added:
        return 0, ["  [SAME] graft fit - no added paragraphs"]

    conn = _conn_positions(old_paras)
    lines, issues = [], 0
    for j, txt, before, after in added:
        if len(txt) < 60:                      # 짧은 삽입(라벨·구절)은 이 축의 대상이 아니다
            continue
        head = txt[:34].replace("\n", " ")
        lines.append(f"    + [{j}] {head}… ({len(txt)}자)")

        # (a) 접속어 자리
        first = _sentences(txt)[0] if _sentences(txt) else txt
        for c in CONNECTIVES:
            if first.startswith(c) and c in conn:
                spots = conn[c]
                if "문단첫" not in spots:
                    issues += 1
                    lines.append(f"      [FAIL] 「{c}」로 문단을 열었는데 host 는 "
                                 f"{len(spots)}회 전부 {'·'.join(sorted(set(spots)))}에 쓴다")

        # (b) 분량 — 이웃 «본문» 문단 분포와
        # ⚠ 창이 «절 경계»를 넘으면 상한이 부풀어 잡아야 할 것을 놓친다(2026-09-11 실측:
        #   고찰 삽입의 이웃에 제한점 545자가 들려 들어와 546자짜리가 통과했다).
        #   제목에서 끊는다 — 제목은 짧고 종결어미가 없다.
        def _is_head(p):
            return len(p) < 60 or not p.rstrip().endswith("다.")

        near = []
        for step, rng in ((-1, range(j - 1, max(-1, j - 1 - window), -1)),
                          (1, range(j + 1, min(len(new_paras), j + 1 + window)))):
            for k in rng:
                if _is_head(new_paras[k]):
                    break                       # 절이 바뀌었다 — 여기서 멈춘다
                if len(new_paras[k]) >= 120:
                    near.append(new_paras[k])
        if len(near) >= 3:
            ln = sorted(len(p) for p in near)
            lo, hi, med = ln[0], ln[-1], ln[len(ln) // 2]
            # 긴 것은 결함이다(그 절을 지배한다). 짧은 것은 대개 의도다 — 내용을 줄인 결과.
            if len(txt) > hi:
                issues += 1
                lines.append(f"      [FAIL] 분량 {len(txt)}자 — 이웃 본문 최장 {hi}자보다 길다"
                             f" (중앙값 {med})")
            elif len(txt) < lo:
                lines.append(f"      [의심] 분량 {len(txt)}자 — 이웃 본문 {lo}~{hi}"
                             f"(중앙값 {med})보다 짧다. 의도한 것이면 그대로")
            else:
                lines.append(f"      [ OK ] 분량 {len(txt)}자 (이웃 {lo}~{hi}, 중앙값 {med})")

    if not lines:
        return 0, ["  [SAME] graft fit - nothing long enough to judge"]
    return issues, [("  [DIFF] graft fit - %d issue(s)" % issues) if issues
                    else "  [ OK ] graft fit"] + lines


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("--render", action="store_true",
                    help="also measure rendered type sizes (Word COM + PyMuPDF)")
    args = ap.parse_args()

    for p in (args.old, args.new):
        if not os.path.exists(p):
            sys.exit(f"missing: {p}")

    print(f"OLD {args.old}\nNEW {args.new}\n")
    tally = {}

    print("1. prose (doc.paragraphs)")
    tally["prose"] = report_seq("prose", prose(args.old), prose(args.new))

    print("\n2. tables (cell text + dimensions)")
    tally["tables"] = report_seq("tables", tables(args.old), tables(args.new))

    print("\n3. geometry (widths, images, margins)")
    tally["geometry"] = report_map("geometry", geometry(args.old), geometry(args.new))

    print("\n5. graft fit (added prose vs host habits)")
    n5, out5 = graft_fit(prose(args.old), prose(args.new))
    tally["graft fit"] = n5
    for l in out5:
        print(l)

    print("\n4. rendered type sizes")
    tally["type sizes"] = None
    if not args.render:
        print("  [SKIP] pass --render to measure (XML w:sz is inherited and unreliable)")
    else:
        try:
            with tempfile.TemporaryDirectory() as td:
                pa, pb = os.path.join(td, "a.pdf"), os.path.join(td, "b.pdf")
                render_pdf(args.old, pa)
                render_pdf(args.new, pb)
                ca, na = type_profile(pa)
                cb, nb = type_profile(pb)
            print(f"      pages {na} -> {nb}")
            print(f"      {'pt':>6} {'OLD':>8} {'NEW':>8}")
            for k in sorted(set(ca) | set(cb)):
                flag = "  <<" if ca.get(k, 0) != cb.get(k, 0) else ""
                print(f"      {k:>6} {ca.get(k,0):>8} {cb.get(k,0):>8}{flag}")
            n_diff = sum(1 for k in set(ca) | set(cb) if ca.get(k, 0) != cb.get(k, 0))
            tally["type sizes"] = n_diff + (na != nb)
            print(f"  [{'DIFF' if tally['type sizes'] else 'SAME'}] type sizes")
        except Exception as e:            # rendering is best-effort, never fatal
            print(f"  [SKIP] render failed: {e}")

    # Summary last, so a filtered or truncated read of this output still cannot
    # miss that an axis changed. "Only X changed" must be read off these counts,
    # never off the body above.
    print("\n" + "=" * 46)
    print("SUMMARY  (changes per axis)")
    for k, v in tally.items():
        print(f"  {k:<12} {'not measured' if v is None else v}")
    print("=" * 46)
    changed = any(v for v in tally.values() if v)
    print("DIFFERENCES FOUND" if changed else "IDENTICAL on all measured axes")
    sys.exit(1 if changed else 0)


if __name__ == "__main__":
    main()
