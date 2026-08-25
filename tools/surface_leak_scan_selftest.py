#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""surface_leak_scan.py의 .docx 지원을 증명한다 (2026-08-11 추가분).

이전에는 .pdf가 아닌 모든 경로를 `open(path, encoding="utf-8")`로 열었다 -- .docx는 zip
바이너리라 즉시 UnicodeDecodeError로 죽었다. agent/tools/audit.py 디스패처가 manuscript/brief
장르에 내놓던 [수동] 힌트가 실제로는 실행하면 깨지는 명령이었다는 뜻이다. 이 셀프테스트는 그
경로가 이제 문단·표 셀 양쪽에서 정상 동작하는지, 그리고 크래시 대신 clean-exit도 여전히
되는지 확인한다.

Run: python surface_leak_scan_selftest.py
"""
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import surface_leak_scan as sls


def make_docx(paragraphs, table_cell=None):
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_cell:
        t = doc.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = table_cell
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def test_does_not_crash_and_finds_paragraph_hit():
    """이전엔 여기서 UnicodeDecodeError로 죽었다."""
    path = make_docx(["이전 버전과 비교하면 차이가 크다.", "정상적인 본문 문장."])
    try:
        hits = sls.scan_file(path, ["이전 버전"])
        assert len(hits) == 1, "문단에서 못 잡음: %r" % hits
        loc, term, snippet = hits[0]
        assert loc == "paragraph 1", "위치 라벨이 다름: %r" % loc
        print("PASS: test_does_not_crash_and_finds_paragraph_hit")
    finally:
        os.unlink(path)


def test_finds_table_cell_hit():
    path = make_docx(["본문에는 없음."], table_cell="사용자가 요청한 값 48,501")
    try:
        hits = sls.scan_file(path, ["사용자가"])
        assert len(hits) == 1, "표 셀에서 못 잡음: %r" % hits
        loc, term, snippet = hits[0]
        assert loc.startswith("table 1 row 1"), "표 위치 라벨이 다름: %r" % loc
        print("PASS: test_finds_table_cell_hit")
    finally:
        os.unlink(path)


def test_clean_docx_exits_empty():
    path = make_docx(["아무 문제 없는 평범한 문장이다."])
    try:
        hits = sls.scan_file(path, ["이전 버전", "사용자가"])
        assert hits == [], "정상 문서에서 오탐: %r" % hits
        print("PASS: test_clean_docx_exits_empty")
    finally:
        os.unlink(path)


if __name__ == "__main__":
    test_does_not_crash_and_finds_paragraph_hit()
    test_finds_table_cell_hit()
    test_clean_docx_exits_empty()
    print("\nALL PASS")
