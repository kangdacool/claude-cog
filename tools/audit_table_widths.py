"""
##################################################################
#####  TABLE WIDTH AUDIT — 숫자가 두 줄로 쪼개지는 사고를 예측  #####
##################################################################

실사고(2026-07-27 학위논문, 같은 유형에 3회): 값 한 글자가 길어져 열 폭을 넘기자
`4,701`이 `4,70` / `1` 로 쪼개졌다. **XML 검사도, 셀↔소스 전수 대조도 전부 통과했다 —
렌더에서만 보였다.** 렌더 확인은 여전히 필요하지만, 이 결함만은 **미리** 계산할 수 있다.

핵심: 진짜 결함은 "칸이 좁다"가 아니라 **끊어지면 안 되는 토큰이 끊긴다**는 것이다.
`1.24 (1.12, 1.37)` 같은 셀은 공백에서 접혀도 되지만 `4,701`이나 `(1.12,`는 안 된다.
그래서 **셀의 가장 긴 공백-구분 토큰**의 렌더 폭을 열 안쪽 폭과 비교한다.

Usage:
  python tools/audit_table_widths.py FILE.docx [--font malgun] [--verbose]
종료 코드: 쪼개질 셀이 있으면 1

⚠️ **국문 셀에서는 과대보고한다(알려진 한계, 미수정).** Word는 한글을 아무 글자에서나 끊으므로
「끊어지면 안 되는 것」은 공백-구분 토큰 전체가 아니라 **그 안의 라틴·숫자 연속구간**이다.
`케냐(국가결핵유병률조사·IMPALA컨소시엄)`은 토큰 전체로 재면 폭이 한참 모자라 보이지만
실제로 보기 싫게 끊기는 건 `IMPALA` 한 구간뿐이다.
→ **판정: 순한글 지적은 무시하고, 라틴 글자나 숫자가 든 것만 고친다**
(`(CASTLE)`이 `(CAST`/`LE)`로, `2024.12`가 `2024.1`/`2`로 갈리는 것이 진짜 결함).
실측 2026-08-20 insaui4 결핵 케이스: 보고 6건 중 손볼 값이 있던 것은 3건.
"""

import argparse
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

EMU_PER_IN = 914400
TWIP_PER_IN = 1440
DEFAULT_CELL_MARGIN_IN = 0.08 * 2      # 좌우 기본 여백 합 (Word 기본 0.08in)

FONTS = {
    "arial": r"C:\Windows\Fonts\arial.ttf",
    "malgun": r"C:\Windows\Fonts\malgun.ttf",
    "times": r"C:\Windows\Fonts\times.ttf",
    "calibri": r"C:\Windows\Fonts\calibri.ttf",
}


def measure(text, ttf, pt, cache={}):
    """텍스트의 렌더 폭(inch). 폰트 파일이 없으면 문자폭 근사로 떨어진다."""
    key = (ttf, pt)
    if key not in cache:
        try:
            from PIL import ImageFont
            cache[key] = ImageFont.truetype(ttf, int(round(pt * 4)))   # 4px/pt로 확대 측정
        except Exception:
            cache[key] = None
    font = cache[key]
    if font is None:                     # 근사: 라틴 0.5em, 한중일 1.0em
        em = pt / 72
        return sum(1.0 if ord(c) > 0x2E80 else 0.5 for c in text) * em
    return font.getlength(text) / (4 * 72)


def col_widths_in(tbl):
    """tblGrid의 gridCol 폭(inch). 없으면 None."""
    from docx.oxml.ns import qn
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return None
    out = []
    for gc in grid.findall(qn("w:gridCol")):
        w = gc.get(qn("w:w"))
        out.append(int(w) / TWIP_PER_IN if w and w.isdigit() else None)
    return out


def cell_font(cell, default_pt, doc_default):
    for p in cell.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                return r.font.size.pt
    return default_pt or doc_default


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--font", default="malgun", choices=sorted(FONTS))
    ap.add_argument("--default-pt", type=float, default=None,
                    help="run에 크기가 없을 때 가정할 pt (기본: 문서 Normal 스타일)")
    ap.add_argument("--verbose", action="store_true")
    a = ap.parse_args()

    from docx import Document
    p = Path(a.file)
    doc = Document(p)
    ttf = FONTS[a.font]
    doc_default = 10.0
    try:
        doc_default = doc.styles["Normal"].font.size.pt or 10.0
    except Exception:
        pass

    findings, checked = [], 0
    for t_i, tbl in enumerate(doc.tables, 1):
        widths = col_widths_in(tbl)
        for r_i, row in enumerate(tbl.rows, 1):
            for c_i, cell in enumerate(row.cells):
                text = cell.text.strip()
                if not text:
                    continue
                # 열 폭: gridCol 우선, 없으면 셀 자체 폭
                w = None
                if widths and c_i < len(widths):
                    w = widths[c_i]
                if w is None and cell.width is not None:
                    w = cell.width / EMU_PER_IN
                if not w:
                    continue
                inner = w - DEFAULT_CELL_MARGIN_IN
                pt = cell_font(cell, a.default_pt, doc_default)
                checked += 1
                # 끊어지면 안 되는 최장 토큰
                tok = max(text.split(), key=len)
                tw = measure(tok, ttf, pt)
                if tw > inner:
                    findings.append((t_i, r_i, c_i + 1, tok, tw, inner, pt))
                elif a.verbose and tw > inner * 0.9:
                    print(f"   near  T{t_i} r{r_i}c{c_i+1}: {tok!r} "
                          f"{tw:.2f}in / {inner:.2f}in")

    print(f"{p.name}: {len(doc.tables)} table(s), {checked} non-empty cells, "
          f"font={a.font} default={doc_default:g}pt")
    for t_i, r_i, c_i, tok, tw, inner, pt in findings:
        print(f"  BREAK  T{t_i} r{r_i}c{c_i}: {tok!r} @{pt:g}pt needs {tw:.2f}in "
              f"but column inner width is {inner:.2f}in — 토큰 중간에서 줄바꿈된다")
    print(f"\n{len(findings)} cell(s) will break inside an unbreakable token")
    if not findings:
        print("표 폭은 안전. 다만 이 검사는 **토큰 쪼개짐만** 본다 — "
              "고아 페이지·각주 넘침은 여전히 렌더해서 눈으로.")
    return 1 if findings else 0


if __name__ == "__main__":
    sys.exit(main())
