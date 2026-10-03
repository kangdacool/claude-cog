# -*- coding: utf-8 -*-
"""office_finalize.py — 사람에게 넘기는 .docx · .pptx · .xlsx 의 «마지막 단계»: 그 Office 프로그램으로 열어 그대로 다시 저장.

    import sys; sys.path.insert(0, "<agent/tools 경로>")
    from office_finalize import finalize
    finalize("out/덱.pptx")                     # 확장자로 Word / PowerPoint / Excel 을 고른다
    python agent/tools/office_finalize.py FILE [FILE ...] [--author "Kang Seo"]

## 왜 있나 (2026-10-02, 실측)

연구자: 「워드에 AI 티는 다 제거했나」 → core 속성(creator)을 고친 python-docx 파일에도 이것이 남아 있었다.
그리고 「앞으로도 word 재저장을 꼭 해라. 모든 연구 파이프라인 산출물 규칙으로」 「powerpoint와 엑셀도 그렇게 해야지」.

    라이브러리     남기는 지문 (기본 템플릿에서 실측)
    python-docx    app.xml Application=「Microsoft Macintosh Word」 14.0 · 단어 수 0 · 템플릿 썸네일
    python-pptx    app.xml 「Microsoft Macintosh PowerPoint」 14.0 · description 「generated using python-pptx」
                   · lastModifiedBy 「Steve Canny」(라이브러리 저자)
    openpyxl       app.xml Application=「Microsoft Excel Compatible / Openpyxl 3.x」

core 칸을 빌더에서 고쳐도 app.xml 은 남는다. 그 프로그램으로 다시 저장하면 이 PC 의 Office 값이 된다.
`audit_doc_properties.py` 가 이 지문을 결함으로 잡는다 — 그래서 재저장을 빼먹으면 첨부 전에 걸린다.

## 지키는 것
- ⛔ `DispatchEx` 만 쓴다(새 프로세스). `Dispatch` 는 사용자가 열어 둔 Word/PowerPoint/Excel 에 붙는다
  — docx_kit.render_docx 머리말의 2026-08-25 사고.
- 저장 «전후 내용 대조»: docx 문단 글자 · pptx 슬라이드별 글자·노트 · xlsx 모든 시트의 셀 값.
  다르면 원본으로 되돌리고 멈춘다(RuntimeError). 재저장이 내용을 바꾼다면 그건 고칠 대상이지 넘길 대상이 아니다.
- 작성자 칸은 `author`(기본 Kang Seo)로, 설명·키워드·마지막 저장자 칸의 도구 흔적은 비운다.
"""
import argparse
import os
import shutil
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass


def _content(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return [p.text for p in d.paragraphs] + cells
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(path).slides):
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append((i, sh.text_frame.text))
                if getattr(sh, "has_table", False) and sh.has_table:
                    out += [(i, c.text) for r in sh.table.rows for c in r.cells]
            if s.has_notes_slide:
                out.append((i, "NOTES", s.notes_slide.notes_text_frame.text))
        return sorted(out, key=str)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=False)
        return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}
    raise ValueError("지원하지 않는 형식: " + ext)


def _props(com_doc, author):
    bp = com_doc.BuiltInDocumentProperties
    for k, v in (("Author", author), ("Comments", ""), ("Keywords", ""), ("Last Author", author)):
        try:
            bp(k).Value = v
        except Exception:
            pass


def _resave(src, dst, author):
    import win32com.client as W
    ext = os.path.splitext(src)[1].lower()
    if ext == ".docx":
        app = W.DispatchEx("Word.Application"); app.Visible = False; app.DisplayAlerts = 0
        try:
            d = app.Documents.Open(src, ReadOnly=True, AddToRecentFiles=False)
            _props(d, author); d.SaveAs2(dst, FileFormat=16, AddToRecentFiles=False); d.Close(False)
        finally:
            app.Quit()
    elif ext == ".pptx":
        app = W.DispatchEx("PowerPoint.Application"); app.DisplayAlerts = 1   # ppAlertsNone
        try:
            d = app.Presentations.Open(src, ReadOnly=True, Untitled=False, WithWindow=False)
            _props(d, author); d.SaveAs(dst, 24); d.Close()                    # 24 = ppSaveAsOpenXMLPresentation
        finally:
            app.Quit()
    elif ext == ".xlsx":
        app = W.DispatchEx("Excel.Application"); app.Visible = False; app.DisplayAlerts = False
        try:
            d = app.Workbooks.Open(src, ReadOnly=True, AddToMru=False)
            _props(d, author); d.SaveAs(dst, 51); d.Close(False)               # 51 = xlOpenXMLWorkbook
        finally:
            app.Quit()
    else:
        raise ValueError("지원하지 않는 형식: " + ext)


## python-docx 기본 템플릿(default.docx)에 든 빈 참고문헌 저장소의 고유 ID. Word 로 다시 저장해도 «그대로 남는다»
## (2026-10-02 실측 — 재저장 뒤 감사 에이전트가 customXml/itemProps1.xml 에서 찾았다. 문자열 검색으로는 안 잡힌다).
## 연구자: 「그렇게 해. 앞으로도 그렇게 하고」 → 재저장 «뒤»에 그 파트를 걷는다. 검사는 audit_doc_properties ⑤.
PYDOCX_TEMPLATE_GUID = "EF278816-EC6F-A645-907D-7F25AECB1D4A"


def strip_template_customxml(path):
    """.docx 에서 python-docx 템플릿 GUID 를 가진 customXml 항목(item·itemProps·rels)을 지우고 그것을 가리키는
    관계·Content_Types 항목도 지운다. 지운 파트 이름 목록을 돌려준다(없으면 빈 목록, 파일은 그대로)."""
    import re
    import zipfile
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    infos = {n: z.getinfo(n) for n in names}
    z.close()
    drop = set()
    for n in names:
        m = re.match(r"customXml/itemProps(\d+)\.xml$", n)
        if m and PYDOCX_TEMPLATE_GUID.encode() in data[n]:
            k = m.group(1)
            drop |= {n, "customXml/item%s.xml" % k, "customXml/_rels/item%s.xml.rels" % k}
    drop &= set(names)
    if not drop:
        return []
    items = [d for d in drop if re.match(r"customXml/item\d+\.xml$", d)]
    for n in names:
        if n.endswith(".rels") and n not in drop:
            t = data[n].decode("utf-8")
            for it in items:
                t = re.sub(r'<Relationship [^>]*Target="[^"]*/%s"[^>]*/>' % re.escape(it), "", t)
            data[n] = t.encode("utf-8")
    ct = data["[Content_Types].xml"].decode("utf-8")
    for d in drop:
        ct = re.sub(r'<Override PartName="/%s"[^>]*/>' % re.escape(d), "", ct)
    data["[Content_Types].xml"] = ct.encode("utf-8")
    tmp = path + ".strip.tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for n in names:
            if n not in drop:
                out.writestr(infos[n], data[n])
    os.replace(tmp, path)
    return sorted(drop)


def finalize(path, author="Kang Seo"):
    """그 Office 프로그램으로 다시 저장하고, 내용이 같을 때만 원본을 바꾼다. 경로를 돌려준다.
    .docx 는 재저장 뒤 python-docx 템플릿의 customXml 흔적을 걷는다(strip_template_customxml)."""
    p = os.path.abspath(path)
    before = _content(p)
    tmpdir = tempfile.mkdtemp(prefix="office_finalize_")
    dst = os.path.join(tmpdir, os.path.basename(p))
    try:
        _resave(p, dst, author)
        if dst.lower().endswith(".docx"):
            strip_template_customxml(dst)
        after = _content(dst)
        if after != before:
            raise RuntimeError("재저장 뒤 내용이 달라졌다 — 원본은 그대로 둔다: " + p)
        shutil.copyfile(dst, p)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    return p


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--author", default="Kang Seo")
    a = ap.parse_args()
    for f in a.files:
        print("재저장:", finalize(f, a.author))
