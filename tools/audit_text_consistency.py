# -*- coding: utf-8 -*-
"""문서의 기계적 오타·표기 일관성 감사 — 맞춤법 검사기가 아니다.

**"기계로 확실히 잡히는 것"만** 전수로 본다. 조사·어미·띄어쓰기 같은 한국어 맞춤법은
여기서 다루지 않는다(검사기 없이는 오탐만 늘린다). 사람 눈이 필요한 영역은 그대로 남긴다.

잡는 것
-------
  괄호·따옴표 짝            열고 안 닫은 것
  참고문헌                  번호 결번 · 미인용 · 존재하지 않는 인용번호
  표기 일관성               단위(ng/mL·kg/m²) · 범위 대시(30-39 vs 30–39) · 인용 범위
                            · p값 공백 · 4자리 수 천단위 콤마
  문장부호                  부호 앞 공백 · 부호 뒤 공백 누락 · 본문 이중 공백
  중복                      영문 중복 단어 · 한글 중복 어절
  고유명사                  분류군 등 지정 용어의 철자 변형

읽는 형식: .hwpx · .docx · .md/.txt  (표는 셀 단위, **문단 경계를 보존**해서 읽는다 —
붙여 읽으면 `US adults,` + `n=43` 이 `US adults,n=43` 이 되어 없는 오타를 만든다)

사용
----
    python audit_text_consistency.py FILE [FILE ...] [--terms 용어1,용어2]

종료코드는 "확실한 결함"이 있을 때만 1. 판단이 필요한 항목은 [확인]으로 따로 낸다.
"""
import argparse
import collections
import os
import re
import sys

# 출력에 em-dash가 섞이는데 Windows 콘솔 기본 코덱(cp949)이 이를 못 찍어 UnicodeEncodeError로
# 스크립트가 죽는다. 2026-08-10에 audit.py로 일괄 실행했을 때 드러났다(deck_audit.py도 같은 버그).
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# 이 부류는 정상인 경우가 많아 실패로 세지 않는다
SOFT = {'문단 끝 종결부호 없음', '본문 이중 공백', '영문 1회 등장 소문자 장단어'}

# 정상 문서에는 나올 수 없는 글자열만. 맞춤법 검사기가 아니므로
# "로써/로서", "다르다/틀리다" 처럼 문맥이 필요한 것은 넣지 않는다(오탐만 는다).
KO_TYPOS = (
    '됬', '됫', '됀', '역활', '오랫만', '있읍', '했읍', '있슴', '갯수', '몇일', '희안',
    '어떡해서', '싯가', '깨끗히', '틈틈히', '일일히', '곰곰히', '번번히', '꼼꼼이',
    '더우기', '아뭏든', '설레임', '금새', '녹슬은', '치뤄', '뒤쳐',
    '유의수즌', '연경군', '통계랑', '회기분석', '유의확율', '표준오류',
)


##################################################################
#####  1. 읽기 — 포맷별로 (본문, 표) 두 갈래로 뽑는다  #####
##################################################################

def read_hwpx(path):
    sys.path.insert(0, os.path.expanduser(
    os.path.join("~", ".claude", "skills", "hwpx-editing", "scripts")))
    import zipfile
    import hwpxlib
    from lxml import etree
    P = hwpxlib.P
    z = zipfile.ZipFile(path)
    body, cells = [], []
    for name in hwpxlib.section_names(z):
        root = etree.fromstring(z.read(name))
        for tbl in root.iter(f"{P}tbl"):
            for row in hwpxlib.table_grid(tbl):       # 문단 경계 보존
                for c in row:
                    if c.strip():
                        cells.append(c)
        for p in root:                                 # 섹션 직계 자식만 = 본문
            if p.tag == f"{P}p" and p.find(f".//{P}tbl") is None:
                t = hwpxlib.own(p)
                if t.strip():
                    body.append(t)
    return body, cells


def read_docx(path):
    from docx import Document
    doc = Document(path)
    body = [p.text for p in doc.paragraphs if p.text.strip()]
    cells = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                # 한 셀 안의 여러 문단도 줄바꿈이다 — hwpx와 같은 규칙
                t = "\n".join(p.text for p in cell.paragraphs)
                if t.strip():
                    cells.append(t)
    return body, cells


# 마크다운에서 «구조»인 줄 — 앞뒤 산문과 이어붙이면 안 된다(표 행을 문장에 붙이는 꼴).
_MD_STRUCT = re.compile(r'^\s*(#{1,6}\s|[-*+]\s|\d+[.)]\s|>\s?|\||---+\s*$|===+\s*$)')
## 목록 항목만 따로 — 이어지는 줄과 붙여 한 단위로 본다(아래 read_plain 참조).
_MD_LIST = re.compile(r'^\s*([-*+]\s|\d+[.)]\s)')


def read_plain(path):
    """.md/.txt를 «문단» 단위로 읽는다 — 줄 단위가 아니라.

    .hwpx/.docx는 문단이 실제로 문단이지만, 마크다운 산문은 부드럽게 줄바꿈된다(soft wrap).
    줄마다 검사하면 여러 줄에 걸친 괄호가 전부 «안 닫힘»으로 잡힌다 -- 2026-08-20에 이
    검사를 파일 저장 훅에 물리자마자 그 오탐이 나왔고, 훅에서 우는 검사는 곧 꺼진다.

    · 빈 줄로 블록을 가르고, 블록 안의 «산문» 줄은 공백으로 이어붙인다.
    · 제목·목록·인용·표 행·구분선은 구조라서 각자 한 단위로 둔다.
    · 코드 펜스(``` 또는 ~~~) 안은 통째로 건너뛴다 -- 코드의 괄호는 산문 규칙을 안 따른다.
    """
    with open(path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    out, buf, in_fence = [], [], False
    def flush():
        if buf:
            out.append(' '.join(buf)); buf.clear()

    for ln in lines:
        if re.match(r'^\s*(```|~~~)', ln):
            flush(); in_fence = not in_fence; continue
        if in_fence:
            continue
        if not ln.strip():
            flush(); continue
        ## 목록 항목은 «시작»만 구조다 — 뒤따르는 들여쓴 줄은 그 항목의 이어진 산문이다.
        ## 한 줄짜리 단위로 확정해 버리면 줄바꿈으로 감긴 괄호가 전부 «안 닫힘»으로 잡힌다
        ## (2026-08-21 SESSION_LOG에서 실제로 2건 오탐). 새 버퍼를 열어 이어붙인다.
        ## 인라인 코드(`...`)도 코드다 — 코드 펜스를 건너뛰는 것과 «같은 원칙»이다.
        ## `측정소코드,측정소명` 같은 열 이름 나열이 「부호 뒤 공백 누락」으로 잡혔다(실사고).
        ## 내용을 지우지 않고 자리표시자로 바꿔 산문 규칙에서만 뺀다.
        ln = re.sub(r'`[^`]+`', '`CODE`', ln)
        if _MD_LIST.match(ln):
            flush(); buf.append(ln.strip()); continue
        if _MD_STRUCT.match(ln):
            flush(); out.append(ln.strip()); continue
        buf.append(ln.strip())
    flush()
    return out, []


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.hwpx':
        return read_hwpx(path)
    if ext == '.docx':
        return read_docx(path)
    if ext in ('.md', '.txt'):
        return read_plain(path)
    raise SystemExit('지원하지 않는 형식: %s (.hwpx/.docx/.md/.txt)' % ext)


##################################################################
#####  2. 검사  #####
##################################################################

##################################################
#####  전문용어 풀어쓰기 · 번역체 규칙표  #####
##################################################

# (표준어 정규식, [풀어쓴 표현들], 이름). 표준어가 같은 문서에 «이미» 있으면 확정 결함으로 적는다.
JARGON_PARAPHRASE = [
    (r'후향', [r'되돌아보(며|아|면서)', r'과거를 돌아보'], '후향적'),
    (r'전향', [r'앞을 향해', r'앞으로 향하는', r'미래를 향해'], '전향적'),
    (r'민감도', [r'얼마나 잘 찾아내는(가|지)를 재'], '민감도'),
    (r'특이도', [r'아닌 것을 아니라고 하는 정도'], '특이도'),
    (r'유병률', [r'병을 갖고 있는 사람의 비율'], '유병률'),
    (r'이중판독', [r'두 번 읽기'], '이중판독'),
    (r'비뚤림|편향', [r'삐뚤어(짐|진)'], '비뚤림'),
]

# 번역체는 «낱말»이 아니라 «짜임»이다. 어휘만 바꾸면 남는다.
TRANSLATIONESE = [
    (r'되어\s?지(다|며|고|는|었)', '이중 피동'),
    (r'[가-힣]히?어\s?지는 것', '이중 피동 명사화'),
    (r'에 있어서', 'in terms of 직역'),
    (r'것으로 사료된다|것으로 보여진다', '번역 상투구'),
    (r'중의 하나(이다|다)', 'one of the 직역'),
    (r'[을를] 가지고 있(다|으며|는)', 'have 직역'),
    (r'[에] 대한 [가-힣]{1,10}[의] [가-힣]{1,10}[의]', 'of 중첩'),
    (r'라고 말할 수 있(다|으며)', 'it can be said 직역'),
    (r'[을를] 통하여|[을를] 통해서 [가-힣]{1,6}[을를] 통해', 'through 남발'),
]

# 세울 필요 없는 것을 세운다 — 「A가 아니라 B」 (2026-09-05)
#
# 규칙 하나다: 참인 것만 단언한다. 부정하려면 부정할 대상을 «먼저 세워야» 하고,
# 세워진 것이 읽는 사람 머리에 남는다. 실측 세 건이 전부 같은 짜임이었다.
#   · 「나는 성적이 하위권이다. 그러니 …」     → 성적을 세웠다 (학생이 읽고 상처받음)
#   · 「실력 문제가 아니라 자료가 안 읽힌 것」  → 실력을 세웠다 (안심시키려다)
#   · 「없어지는 건 아니고 그대로 남아요」      → 없어질 가능성을 세웠다
# 셋째는 그 교훈을 메모리에 적은 «다음 문장»에서 났다. 사례를 쌓아서는 안 고쳐진다.
#
# ⚠ 「A가 아니라 B」가 늘 잘못은 아니다. A 가 이미 무대에 있으면 정당한 대조다
#   (「등급이 아니라 생략 금지 표시다」 — 등급은 그 문서의 주제어다).
#   그래서 «부정당한 말이 문서 다른 곳에도 있는가»를 함께 본다. 없으면 우리가 세운 것이다.
#   JARGON_PARAPHRASE 가 「표준어가 같은 문서에 있는가」를 보는 것과 같은 방식이다.
STRAWMAN = [
    r'([가-힣A-Za-z]{2,12})\s*문제가\s*아니(?:라|고|며)',
    r'([가-힣]{2,12})(?:는|은|이|가)\s*아니(?:라|고)\s',
    r'([가-힣]{2,12})\s*건\s*아니(?:고|라)',
]

# 읽는 사람을 규정하는 전제. 처방만 남기면 지시는 그대로 작동한다.
PERSON_PREMISE = [
    (r'(?:나는|저는|학생은|독자는)\s*[^\n]{0,12}'
     r'(?:하위권|상위권|성적이|실력이|초보|잘\s*못하|부족하)', '사람을 규정하는 전제'),
]


def audit(body, cells, terms=(), rendered=True):
    allt = body + cells
    hits = collections.OrderedDict()

    def add(cat, loc, detail):
        hits.setdefault(cat, []).append((loc, detail))

    # -- 괄호 짝: **열고 안 닫은 것만** 본다 --
    # 남는 닫는 괄호는 거의 항상 열거("1) 중증 …", "가) …")라 오탐이 된다.
    # 어포스트로피(Faith’s)와 여는 따옴표 없는 인용도 같은 이유로 제외.
    for i, t in enumerate(allt):
        for a, b in (('(', ')'), ('[', ']'), ('“', '”')):
            depth = 0
            unclosed = 0
            for ch in t:
                if ch == a:
                    depth += 1
                elif ch == b:
                    depth = max(0, depth - 1)
            unclosed = depth
            if unclosed:
                add('괄호 열고 안 닫음', '#%d' % i,
                    '%s%s %d개 미종결 | %s' % (a, b, unclosed, t[:70]))

    # -- 참고문헌 번호·인용 커버리지 --
    refs = {}
    for t in body:
        # 저자명 첫 글자는 라틴 대문자 **또는 한글**이다. [A-Z]만 보면 국문 문헌
        # (예: "9. 교육부. 교권보호위원회 ...")이 통째로 안 잡혀, 정상적으로 인용된
        # [9]·[10]이 "존재하지 않는 인용번호"로 보고된다 — 2026-08-10 실제 오탐.
        m = re.match(r'^\s*(\d{1,3})\.\s+[A-Z가-힣]', t)
        # 국문 양식의 «중간제목»(1. 자료와 대상)이 여기 걸린다. 참고문헌 한 항목은
        # 저자·제목·지면을 담아 길고 문장부호가 여럿이지만 절 제목은 짧다 -- 그 차이로
        # 가른다(2026-09-04 실측: 절 제목 6개가 통째로 «미인용 참고문헌»이 됐다).
        if m and (len(t) < 30 and sum(t.count(c) for c in ".,;:()") <= 1):
            m = None
        if m:
            refs[int(m.group(1))] = t[:60]
    if len(refs) >= 5:                                   # 참고문헌 목록이 있는 문서만
        miss = [i for i in range(1, max(refs) + 1) if i not in refs]
        if miss:
            add('참고문헌 번호 결번', 'refs', str(miss))
        # 인용 마커는 «두 가지» 형식이 있다: 대괄호 [1,2] 와 위첨자 <sup>1,2</sup>.
        # 대괄호만 보면 위첨자 원고(AJE가 요구하는 형식)에서 «전 참고문헌이 미인용»
        # 으로 보고되고, 아래 순서 검사도 통째로 무력화된다 — 2026-08-19 실제 오탐.
        CITE_RE = re.compile(r'\[([0-9,\-–\s]+)\]|<sup>([0-9,\-–\s]+)</sup>')

        def cited_nums(t):
            out = []
            for m in CITE_RE.finditer(t):
                inner = m.group(1) or m.group(2)
                for part in inner.split(','):
                    part = part.strip()
                    if re.search(r'[-–]', part):
                        a2, b2 = re.split(r'[-–]', part)[:2]
                        if a2.strip().isdigit() and b2.strip().isdigit():
                            out.append((m.start(), list(range(int(a2), int(b2) + 1))))
                    elif part.isdigit():
                        out.append((m.start(), [int(part)]))
            return out

        cited = set()
        for t in allt:
            if re.match(r'^\s*\d{1,3}\.\s+[A-Z가-힣]', t):
                continue
            for _, ns in cited_nums(t):
                cited.update(ns)
        never = sorted(set(refs) - cited)
        ghost = sorted(x for x in cited - set(refs) if x <= max(refs) + 50)
        if never:
            add('미인용 참고문헌', 'refs', str(never))
        if ghost:
            add('존재하지 않는 인용번호', 'refs', str(ghost))
        add('[확인] 참고문헌 개수', 'refs', '%d개 (1~%d)' % (len(refs), max(refs)))

        # -- 첫인용 순서 (Vancouver/AMA) --
        # 존재·양방향 대응이 다 맞아도 «순서»는 따로다. 실제 사고(2026-08-19,
        # 한 원고에서): cross-verify 44/44 통과 + "1~52 전부 본문 인용·목록 존재"
        # 확인까지 한 원고를 지도교수께 보냈는데, 첫인용 순서가
        # 1 2 6 8 9 7 ... 17 46 47 48 49 18 3 4 5 ... 로 어긋나 있었다.
        # 같은 검사를 자매 프로젝트에서는 «했는데» 이쪽에 안 돌린 것이 원인이다.
        first_at = {}
        for pos, t in enumerate(allt):
            if re.match(r'^\s*\d{1,3}\.\s+[A-Z가-힣]', t):
                continue                                  # 목록 줄은 본문이 아니다
            for off, ns in cited_nums(t):
                for n in ns:
                    first_at.setdefault(n, (pos, off))
        seq = sorted(first_at, key=lambda n: first_at[n])
        if seq and seq != sorted(seq):
            bad_at = next((i for i in range(1, len(seq)) if seq[i] < seq[i - 1]), None)
            detail = '첫인용 순서: %s' % ' '.join(str(x) for x in seq[:14])
            if bad_at is not None:
                detail += ' ... (첫 위반: %d 다음에 %d)' % (seq[bad_at - 1], seq[bad_at])
            # 목록이 «저자 알파벳순»이면 번호가 첫인용 순서가 아닌 것이 의도일 수
            # 있다(알파벳-번호 병용 스타일). 그 경우만 판단 항목으로 낮춘다.
            heads = [refs[k].split('.', 1)[-1].strip()[:20] for k in sorted(refs)]
            alpha = heads == sorted(heads, key=str.lower)
            if alpha:
                add('[확인] 참고문헌 번호가 첫인용 순서가 아님', 'refs',
                    detail + ' | 목록이 알파벳순이므로 의도일 수 있음')
            else:
                add('참고문헌 번호가 첫인용 순서가 아님', 'refs', detail)

    # -- 저자-연도 인용의 «말이 되는가» 검사 --
    # 번호 인용을 저자-연도로 일괄 변환하는 스크립트는 «예외 없이 끝나고 개수도 맞으면서»
    # 조용히 망가진다. 2026-08-19에 실제로 두 종류가 원고에 실려 나갔다:
    #   - 'van der Laan MJ' 의 성을 첫 단어 'van' 으로 잡아 (van & Rose, 2011)
    #   - 'J Clin Diagn Res. 2013;7(9):2063-2067' 에서 «끝 페이지»를 연도로 잡아
    #     (Dable et al., 2067)
    # 둘 다 개수 검사·예외 검사를 통과했고, 사람이 «읽어서» 발견했다. 그래서 여기서
    # 기계로 잡는다 — 변환기마다 다시 짜지 말고 산출물 쪽에서 본다.
    PARTICLES = {'van', 'von', 'de', 'der', 'den', 'del', 'della', 'di', 'da', 'du',
                 'la', 'le', 'les', 'el', 'al', 'bin', 'ibn', 'ter', 'ten'}
    # 다중 인용 '(A et al., 2011; B & C, 2007)' 은 세미콜론으로 «먼저» 쪼갠다.
    # 통째로 정규식을 걸면 마지막 하나만 보게 되어, 앞쪽에 있는 깨진 인용을 놓친다.
    paren = re.compile(r'\(([^()]{4,300})\)')
    one = re.compile(r'^(.+?),\s*((?:1[6-9]|20)?\d{2,4})([a-z])?$')
    seen_ad = 0
    for t in allt:
        for pm in paren.finditer(t):
            for part in pm.group(1).split(';'):
                part = part.strip()
                m = one.match(part)
                if not m or not re.search(r'[A-Za-zÀ-ÿ]', m.group(1)):
                    continue
                who, ys = m.group(1).strip(), m.group(2)
                if not ys.isdigit() or len(ys) != 4:
                    continue
                seen_ad += 1
                yr = int(ys)
                if not (1600 <= yr <= 2035):
                    add('저자-연도 인용의 연도가 비현실적', 'cite',
                        '(%s) | 페이지 번호를 연도로 잡았을 수 있음' % part)
                head = re.split(r'\s*(?:&|,)\s*', who)[0].strip()
                toks = head.split()
                if toks and toks[0].lower() in PARTICLES and len(toks) == 1:
                    add('저자-연도 인용의 성이 전치사 한 단어', 'cite',
                        "(%s) | 'van der Laan'을 'van'으로 자른 형태일 수 있음" % part)
    if seen_ad:
        add('[확인] 저자-연도 인용 개수', 'cite', '%d개' % seen_ad)

    # -- 표기 일관성: 한 문서 안에서 두 가지 형태가 섞이면 결함 --
    def mixed(cat, pattern, key=lambda m: m.group(0), scope=None, pre=None):
        c = collections.Counter()
        for t in (scope if scope is not None else allt):
            for m in re.finditer(pattern, pre(t) if pre else t):
                c[key(m)] += 1
        if len(c) > 1:
            add(cat, 'all', dict(c))

    # ISO 날짜(2022-01-01)의 '-'는 범위 대시가 아니라 «날짜 구분자»다.
    # 지우지 말고 «마스킹»한다 — 지우면 앞뒤 숫자가 붙어 없던 범위가 생긴다.
    #
    # ⭐ 2026-09-08 — «식별자»도 같은 이유로 마스킹한다.
    #    실측: 참고문헌 49개짜리 투고 원고가 이 검사에서 «항상» 실패했는데, 하이픈 히트가
    #    전부 DOI(`10.1186/1476-072X-13-38`)·ORCID(`0009-0009-6082-2368`)·UCI 였다.
    #    진짜 숫자 범위는 100% en-dash 로 정확했다. 즉 **원고를 고칠 방법이 없는 실패**다.
    #    ⛔ 늘 실패하는 검사는 읽히지 않고, 그러면 «진짜» 실패가 그 뒤에 묻힌다
    #       (같은 파일이 이미 아는 교훈: 시끄러운 검사는 조용한 검사와 결과가 같다).
    ID_PATTERNS = (
        r'\d{4}-\d{2}-\d{2}',                      # ISO 날짜
        r'\b10\.\d{4,9}/\S+',                       # DOI
        r'\bdoi:\s*\S+',                            # doi: 접두
        r'\b\d{4}-\d{4}-\d{4}-\d{3}[\dXx]\b',       # ORCID
        r'\bG?7?0?4?-?\bUCI\s*\S+',                 # UCI (표기 흔들림 대비)
        r'\bISSN\s*\d{4}-\d{3}[\dXx]\b',
        # 전화·팩스 — 「범위」가 아니라 식별자다. 2026-09-29 실측: 원고 겉표지의
        # `Tel: +82-10-2739-2458` 하나 때문에 이 검사가 «매 빌드마다» 떴다(260921·260929 동일).
        r'\+\d{1,3}[-\d]{6,}',
        r'(?:Tel|TEL|Fax|FAX|전화|팩스)\s*[:：]?\s*\+?[\d\-]{7,}',
    )

    def iso_mask(t):
        for pat in ID_PATTERNS:
            t = re.sub(pat, lambda m: m.group(0).replace('-', '␣'), t)
        return t

    mixed('단위 표기 혼용 (ng/mL)', r'ng/m[lL]')
    mixed('단위 표기 혼용 (BMI)', r'kg/m[2²]')
    # ⚠ 이 검사는 이름보다 «좁다»: 두 자리 정수 범위(18-34)만 본다. 쉼표·소수점이 낀 범위
    #    (18,300-34,300 · 2.83-3.70 · 0.77-1.22)는 전후 자릿수 조건에 걸려 **잡히지 않는다**
    #    (2026-09-29 확인). 넓히면 좌표·용량·모델명까지 걸리므로 그대로 두되, 좁다는 것을 알고 쓴다.
    mixed('숫자 범위 대시 혼용', r'(?<!\d)\d{2}\s*([-–])\s*\d{2}(?!\d)',
          lambda m: m.group(1), pre=iso_mask)
    mixed('인용 범위 대시 혼용', r'\[\d{1,3}([-–])\d{1,3}\]', lambda m: m.group(1))
    mixed('p값 공백 혼용', r'p(\s*)[<=>](\s*)0',
          lambda m: '공백 있음' if (m.group(1) or m.group(2)) else '공백 없음', body)
    # 연도(19xx·20xx)는 콤마를 안 찍는 게 정상이므로 제외한다.
    # ':' 바로 뒤 4자리도 제외 — 수량이 아니라 «코드»다(EPSG:5179, 포트 등). 콤마를 찍으면 틀린다.
    mixed('4자리 수 천단위 콤마 혼용',
          r'(?<![\d,.:])(\d,\d{3}|(?!19\d\d|20\d\d)\d{4})(?![\d,.%])',
          lambda m: '콤마' if ',' in m.group(1) else '없음', cells)

    # -- 마크다운 잔재 (생성된 문서에 «서식 기호가 글자로» 찍힌 것) --
    # docx/hwpx 빌더는 마크다운을 해석하지 않는다. 초안을 md로 쓰다 그대로 넘기면 별표·백틱이
    # 본문에 남는다 — 렌더해서 눈으로 보기 전엔 안 보이고, 숫자 검증도 통과한다.
    # ⚠ 좁게 잡는다: 유의성 별표(1.23**)·각주 표시(*)는 «짝»이 아니므로 안 걸린다.
    # ⚠⚠ **.md/.txt에는 돌리지 않는다** — 거기서 `**굵게**`는 잔재가 아니라 «의도된 서식»이다.
    #     처음엔 형식을 안 가려서 SESSION_LOG.md 하나에 400건을 신고했다(2026-08-21).
    for i, t in enumerate((body + cells) if rendered else []):
        where = '#%d' % i
        for m in re.finditer(r'\*\*(?=\S)([^*\n]{2,120}?)(?<=\S)\*\*', t):
            add('마크다운 굵게가 글자로 남음', where, m.group(0)[:60])
        for m in re.finditer(r'`(?=\S)([^`\n]{1,80}?)(?<=\S)`', t):
            add('마크다운 코드표시가 글자로 남음', where, m.group(0)[:60])
        # 작업노트 표기 누출 — «»·⭐·🔴 는 agent/feedback·MEMORY·SESSION_LOG 의 강조 표기다.
        # 그 문체로 오래 쓰다 보면 산출물 산문에 그대로 딸려 나온다(2026-09-03: 학생 배포
        # 문서군에 26건). 마크다운 잔재와 같은 부류 — «서식 기호가 글자로 남은 것»이다.
        for m in re.finditer(r'«[^»\n]{1,80}»', t):
            add('작업노트 표기가 산출물에 남음(이중꺽쇠)', where, m.group(0)[:60])
        for m in re.finditer(r'[⭐🔴🟡🟢⚠✅❌★]', t):
            add('작업노트 기호가 산출물에 남음', where,
                t[max(0, m.start() - 20):m.start() + 25])

        # 전문용어를 «고유어로 풀어쓴 것». 2026-09-03 실측: 「AI 말투 같다」는 지적을
        # 「전문용어를 풀어써라」로 잘못 받아 후향적/전향적을 「되돌아보며」·「앞을 향해
        # 관찰하여」로 바꿨다. 학술 문서에서 이건 AI 티가 아니라 «비전문가 티»다.
        # ⚠ 같은 문서 다른 곳에는 「후향 27편·전향 3편」이 멀쩡히 있었다 — 그래서
        #   이 검사는 «표준어가 문서 안에 이미 있는가»를 함께 본다(있으면 확정 결함).
        _whole = ' '.join(body + cells)
        for std, pats, label in JARGON_PARAPHRASE:
            for pat in pats:
                for m in re.finditer(pat, t):
                    seen_std = re.search(std, _whole) is not None
                    add('전문용어를 풀어썼다(%s%s)' % (label, ' — 표준어가 같은 문서에 있다'
                                                if seen_std else ''),
                        where, m.group(0)[:60])

        # 번역체의 «구조적» 지문. 낱말이 아니라 짜임이라 어휘를 바꿔도 남는다.
        for pat, label in TRANSLATIONESE:
            for m in re.finditer(pat, t):
                add('번역체 어투(%s)' % label, where,
                    t[max(0, m.start() - 18):m.start() + 28])

    # -- 문장부호·공백·중복 (본문만; 표는 줄바꿈이 많아 오탐) --
    for i, t in enumerate(body):
        # 저널 약어는 같은 낱말이 겹치는 것이 «정상»이다 -- "Teach Teach Educ"
        # (Teaching and Teacher Education), "Nurs Nurs Res" 등. 참고문헌 목록에서 반복
        # 오탐이 났다(2026-08-22 teacher/paper).
        # ⚠ 「대문자로 시작 + 뒤에 약어」만으로 거르면 "The The Ministry"가 새 나간다(실측).
        #    그래서 그 문단이 «서지»임을 따로 확인한다 -- 연도;권 또는 doi:/PMID 가 있어야 한다.
        is_citation = re.search(r'\b(19|20)\d\d;\d|doi:|PMID\s*\d', t) is not None
        for m in re.finditer(r'\b([A-Za-z]{3,})\s+\1\b', t):
            if (is_citation and m.group(1)[0].isupper()
                    and re.match(r'\s+[A-Z][a-z]{2,9}\b', t[m.end():])):
                continue
            add('영문 중복 단어', '#%d' % i, m.group(0))
        for m in re.finditer(r'([가-힣]{2,})\s+\1(?![가-힣])', t):
            add('한글 중복 어절', '#%d' % i, m.group(0))
        for m in re.finditer(r'\s+[,.](?=\s|$)', t):
            add('문장부호 앞 공백', '#%d' % i, repr(t[max(0, m.start() - 20):m.end() + 12]))
        ## 파일명의 확장자 점은 문장부호가 아니다 — 'README_품질.md는'은 올바른 한국어다.
        ## 마스킹해서 검사에서 빼되, 문장 끝 점("마쳤다.다음은")은 그대로 걸리게 둔다.
        ## 확장자는 «소문자»로 한정한다 — "크다.The"처럼 대문자로 시작하는 영문 문장은
        ## 여전히 잡아야 하기 때문이다.
        t_fn = re.sub(r'\S+\.[a-z0-9]{1,5}(?![a-zA-Z0-9])',
                      lambda m: m.group(0).replace('.', '․'), t)
        for m in re.finditer(r'[가-힣][,.][가-힣A-Za-z]', t_fn):
            add('부호 뒤 공백 누락', '#%d' % i, repr(t_fn[max(0, m.start() - 15):m.end() + 12]))
        for m in re.finditer(r'[가-힣A-Za-z0-9)\]] {2,}[가-힣A-Za-z0-9(\[]', t):
            add('본문 이중 공백', '#%d' % i, repr(m.group(0)))

    # -- 세울 필요 없는 것을 세운다 --
    # ⚠ 렌더된 산출물 전용 구간에 두면 안 된다. 이건 «초안부터» 걸려야 하는 규칙이다
    #   (2026-09-05: 처음에 rendered 루프 안에 넣었다가 자기시험에서 전건 무통과로 잡혔다).
    _all = ' '.join(allt)
    for i, t in enumerate(allt):
        for pat in STRAWMAN:
            for m in re.finditer(pat, t):
                word = m.group(1)
                # 부정당한 말이 문서 다른 곳에도 있으면 «이미 무대에 있던 것»이라 정당한 대조다.
                if len(re.findall(re.escape(word), _all)) <= 1:
                    add('세운 적 없는 것을 부정한다(참인 것만 단언할 것)', '#%d' % i,
                        t[max(0, m.start() - 15):m.end() + 20])
        for pat, label in PERSON_PREMISE:
            for m in re.finditer(pat, t):
                add('%s — 처방만 남길 것' % label, '#%d' % i, m.group(0)[:60])

    # -- 자주 틀리는 한국어 표기 (사전 아님 — 오타로만 나오는 형태만) --
    # 맞춤법 검사가 아니다. "정상 문서에는 나올 수 없는 글자열"만 담는다.
    for i, t in enumerate(allt):
        for w in KO_TYPOS:
            if w in t:
                add('한국어 오표기', '#%d' % i, '%s | %s' % (w, t[:60]))

    # -- 지정 용어 철자 변형 --
    blob = '\n'.join(allt)
    for term in terms:
        if not term:
            continue
        stem = term[:max(4, len(term) - 4)]
        # 한글 조사가 붙은 형태(Firmicutes가)는 변형이 아니므로 라틴 문자만 이어 붙인다
        var = {v for v in re.findall(r'\b%s[A-Za-z]*' % re.escape(stem), blob) if v != term}
        # 실제 다른 단어(하위 분류군 등)는 접두사가 같아도 정상이므로 확인 항목으로만
        if var:
            add('[확인] 용어 철자 변형', term, sorted(var)[:8])
    return hits


##################################################################
#####  3. 출력  #####
##################################################################

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='+')
    ap.add_argument('--terms', default='',
                    help='철자 변형을 확인할 고유명사(쉼표 구분)')
    args = ap.parse_args()
    terms = [t.strip() for t in args.terms.split(',') if t.strip()]

    hard = 0
    for path in args.files:
        body, cells = read_any(path)
        print('\n===== %s  (본문 %d문단 · 표 셀 %d개) =====' % (
            os.path.basename(path), len(body), len(cells)))
        ## «렌더된 산출물»인가(.docx/.hwpx) 아니면 «원문»인가(.md/.txt).
        ## 마크다운 잔재 검사는 앞의 것에만 뜻이 있다.
        rendered = os.path.splitext(path)[1].lower() in ('.docx', '.hwpx')
        hits = audit(body, cells, terms, rendered=rendered)
        if not hits:
            print('   결함 없음')
            continue
        for cat, v in hits.items():
            soft = cat.startswith('[확인]') or cat in SOFT
            print('\n  %s %s — %d건' % ('[확인]' if soft else '[결함]',
                                        cat.replace('[확인] ', ''), len(v)))
            seen = set()
            for loc, d in v:
                k = (loc, str(d))
                if k in seen:
                    continue
                seen.add(k)
                print('     %-8s %s' % (loc, d))
            if not soft:
                hard += len(v)
    print('\n확실한 결함 %d건 (그 외 [확인] 항목은 눈으로 판단)' % hard)
    return 1 if hard else 0


if __name__ == '__main__':
    raise SystemExit(main())
