"""
##################################################################
#####  FONT SIZE AUDIT — 만성 소글씨 편향을 검사로              #####
##################################################################

CORE §2①: "글씨를 항상 너무 작게 만든다. 사고 5회, '너무 크다'는 피드백은 단 한 번도 없었다."
산문 규칙으로는 반복해서 깨졌으므로 검사로 옮긴다.

세 가지를 본다 — 셋 다 다른 경로로 글씨를 작게 만든다:

  [RUN]   pptx/docx 안의 실제 글자 크기가 바닥값 미만
  [SCALE] 이미지를 native보다 작게 배치 → **그림 안에 구워진 글씨가 그 비율만큼 줄어든다.**
          11pt로 그린 축 라벨을 0.7배로 넣으면 7.7pt가 된다. XML 검사로도 숫자 대조로도
          안 잡히고, 렌더에서만 보인다.
          ⚠ **native 는 PNG 의 DPI 메타데이터에서 나온다 — 그 값이 «거짓»일 수 있다.**
          2026-09-08 실측: 논문 PDF 에서 꺼낸 출판사 그림이 실제로는 300 dpi(인쇄폭 7.00in)
          인데 메타데이터에 96 이 박혀 있었다. 그러면 native 가 21.9in 으로 읽혀 **어떤
          슬라이드에 놓아도** 축소 배치로 신고된다 — 그림을 줄인 적이 없는데도.
          → SCALE 이 «그 파일의 그림 전부»에 뜨면 먼저 `PIL.Image.open(p).info['dpi']` 를
          확인한다. PDF 에서 꺼낸 그림이면 `page.get_image_rects(xref)` 로 실제 인쇄 폭을
          재서 `im.save(out, dpi=(d, d))` 로 바로잡은 뒤 다시 볼 것.
          «한두 장만» 뜨면 그건 진짜 축소 배치다 — 글씨를 줄이지 말고 그림을 키운다.
          단 «설명용 그림»(장비 사진·부위 지도)은 대체 텍스트를 「설명용 …」으로 시작하게
          두면 바닥이 0.60 으로 낮아진다(pptx 만 — ILLUSTRATION_FLOOR 주석).
  [SRC]   R/Python 플롯 소스의 크기 리터럴(cex=, size=, fontsize=)이 바닥값 미만

Usage:
  python tools/audit_font_sizes.py FILE...            # .pptx .docx .R .py
  python tools/audit_font_sizes.py deck.pptx --profile slide-ko
  python tools/audit_font_sizes.py fig.R --min-src 10

프로파일 (pt) — **확장자가 아니라 장르로 고른다.** .pptx라고 다 slide가 아니다:
  slide-ko  본문 16 / 제목 30   한글 슬라이드 (한글은 라틴보다 조밀하다)
  slide-en  본문 14 / 제목 24
  poster    본문 24 / 제목 42   학술대회 포스터 — 벽에서 읽는다
  doc       본문 9  / 제목 11   저널 표·캡션
종료 코드: 위반 시 1

같이 볼 것: `claude-config/skills/pptx-editing/scripts/poster_kit.py`의 `min_font_report()`가
[RUN] 검사를 독립적으로(이 파일을 import하지 않고) 다시 구현하고 있다 — 공개 배포되는 스킬이라
이 사설 도구에 의존할 수 없기 때문이다. **[RUN]의 판정 로직을 바꾸면 그쪽도 봐야 한다.**
장르를 자동으로 고르는 것은 `agent/tools/audit.py`가 한다 — 이 스크립트를 직접 부를 땐 --profile을 꼭 줄 것.
"""

import argparse
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PROFILES = {
    "slide-ko": dict(body=16, title=30, caption=13),
    "slide-en": dict(body=14, title=24, caption=10),
    "doc": dict(body=9, title=11, caption=8),
    # 학술대회 포스터는 벽에서 읽는다. slide-ko(본문 16)로 검사하면 18pt 포스터가
    # 그냥 통과한다 — 2026-08-11에 이 프로파일이 없어 실제로 그렇게 돌고 있었다.
    # 바닥값 24는 poster_rules.md의 규칙(kf()의 assert와 같은 값)이다.
    "poster": dict(body=24, title=42, caption=24),
}
SCALE_FLOOR = 0.95      # native 대비 이 아래로 줄여 배치하면 구워진 글씨가 줄어든다
## «설명용 그림»(장비 사진·측정 부위 지도처럼 청중이 글씨를 «읽지» 않는 그림)은 바닥을 낮춘다.
## 2026-10-01 한 코호트 자료 소개 덱 13장: 연구자가 직접 키워 놓은 fNIRS 장비·ROI 그림이 0.94×·0.87× 로
## FAIL 이 났다 — 데이터 그림의 규칙을 삽화에 그대로 쓴 것이다. 그렇다고 검사에서 빼면 같은 표시로
## 데이터 그림이 숨는다. 그래서 «면제»가 아니라 «바닥을 낮춤»이고, 표시는 그림의 대체 텍스트(alt text,
## p:cNvPr/@descr)가 「설명용」 또는 「illustration」으로 «시작»할 때만 인정한다 — 빌더가 일부러 써야 붙는다.
ILLUSTRATION_FLOOR = 0.60
_ILLUS_RE = re.compile(r"^\s*\[?\s*(설명용|illustration)", re.I)


def _is_illustration(sh):
    try:
        return bool(_ILLUS_RE.match(sh._element.nvPicPr.cNvPr.get("descr") or ""))
    except Exception:
        return False
SRC_FLOOR = 9           # 플롯 소스의 크기 리터럴 바닥값


def audit_pptx(path, floors):
    from pptx import Presentation
    from pptx.util import Emu
    out = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == 13 and sh.image is not None:      # PICTURE
                try:
                    nat_w = sh.image.size[0] / (sh.image.dpi[0] or 300)
                    placed = Emu(sh.width).inches
                    illus = _is_illustration(sh)
                    floor = ILLUSTRATION_FLOOR if illus else SCALE_FLOOR
                    if nat_w and placed / nat_w < floor:
                        out.append((f"slide {i}", "SCALE", f"{sh.name!r} placed at "
                                    f"{placed / nat_w:.2f}× native"
                                    + (" (설명용 그림 바닥 0.60 도 밑돈다)" if illus
                                       else " — 구워진 글씨가 그만큼 줄어든다")))
                except Exception:
                    pass
            # ⚠ 2026-09-30: run 크기만 보던 판은 크기를 «문단 기본값(pPr/defRPr)»에 두는 빌더(kit_mono)의
            #   글자를 거의 다 건너뛰었고, 표 셀은 아예 안 봤다 — 12pt 참고문헌이 13pt 바닥값을 «통과»했다.
            #   run 크기가 없으면 문단 기본값을 쓰고, 표 셀도 같은 규칙으로 본다.
            frames = []
            if sh.has_text_frame:
                frames.append(sh.text_frame)
            if getattr(sh, "has_table", False):
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
            for p in (p for tf in frames for p in tf.paragraphs):
                for r in p.runs:
                    size = r.font.size or p.font.size
                    if size is None or not r.text.strip():
                        continue
                    pt = size.pt
                    # 두 층으로 나눈다. 캡션 바닥값 미만 = 위반. 그 사이 = 캡션이라면 의도일
                    # 수 있으므로 경고만 — 한 층으로 재면 의도한 캡션까지 전부 신고해
                    # 목록이 신뢰를 잃는다.
                    if pt < floors["caption"]:
                        out.append((f"slide {i}", "RUN",
                                    f"{pt:g}pt < {floors['caption']}pt — {sh.name!r}: "
                                    f"{r.text.strip()[:45]!r}"))
                    elif pt < floors["body"]:
                        out.append((f"slide {i}", "warn",
                                    f"{pt:g}pt < body {floors['body']}pt — {sh.name!r}: "
                                    f"{r.text.strip()[:40]!r} (캡션이면 의도일 수 있음)"))
    return out


def audit_docx(path, floors):
    from docx import Document
    from docx.shared import Emu
    out = []
    doc = Document(path)
    seen = set()
    for para in doc.paragraphs:
        for r in para.runs:
            if r.font.size is None:
                continue
            pt = r.font.size.pt
            if pt < floors["caption"] and (pt, r.text[:20]) not in seen:
                seen.add((pt, r.text[:20]))
                out.append(("body", "RUN", f"{pt:g}pt < {floors['caption']}pt — "
                                           f"{r.text.strip()[:45]!r}"))
    for t_i, tbl in enumerate(doc.tables, 1):
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for r in para.runs:
                        if r.font.size is not None and r.font.size.pt < floors["caption"]:
                            key = (t_i, r.font.size.pt)
                            if key in seen:
                                continue
                            seen.add(key)
                            out.append((f"table {t_i}", "RUN",
                                        f"{r.font.size.pt:g}pt < {floors['caption']}pt"))
    for shape in doc.inline_shapes:
        try:
            placed = Emu(shape.width).inches
            # python-docx는 native px를 직접 안 준다 — 관계에서 이미지 바이트를 읽는다
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            part = doc.part.related_parts[rid]
            from PIL import Image
            import io
            im = Image.open(io.BytesIO(part.blob))
            dpi = (im.info.get("dpi") or (300, 300))[0] or 300
            nat = im.size[0] / dpi
            if nat and placed / nat < SCALE_FLOOR:
                out.append(("body", "SCALE",
                            f"image placed at {placed / nat:.2f}× native "
                            f"({placed:.2f}in vs {nat:.2f}in)"))
        except Exception:
            pass
    return out


SRC_PATTERNS = [
    (re.compile(r"\bcex(?:\.\w+)?\s*=\s*([0-9.]+)"), "cex", 0.9),       # R base: cex는 배율
    (re.compile(r"\bsize\s*=\s*([0-9.]+)"), "size", None),               # ggplot element_text
    (re.compile(r"\bfontsize\s*=\s*([0-9.]+)"), "fontsize", None),       # matplotlib
    (re.compile(r"\bfont_size\s*=\s*([0-9.]+)"), "font_size", None),
    (re.compile(r'"font\.size"\s*:\s*([0-9.]+)'), "font.size", None),
]


def audit_source(path, min_src):
    out = []
    for ln, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
        if line.lstrip().startswith("#"):
            continue
        for pat, label, ratio_floor in SRC_PATTERNS:
            for m in pat.finditer(line):
                v = float(m.group(1))
                if ratio_floor is not None:      # cex는 배율이라 별도 바닥값
                    if v < ratio_floor:
                        out.append((f"line {ln}", "SRC", f"{label}={v:g} < {ratio_floor} (배율)"))
                elif v < min_src:
                    out.append((f"line {ln}", "SRC", f"{label}={v:g} < {min_src}pt"))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--profile", default="slide-ko", choices=sorted(PROFILES))
    ap.add_argument("--min-src", type=float, default=SRC_FLOOR)
    a = ap.parse_args()
    floors = PROFILES[a.profile]

    total = 0
    for f in a.files:
        p = Path(f)
        if not p.exists():
            print(f"!! not found: {p}")
            continue
        ext = p.suffix.lower()
        if ext == ".pptx":
            hits = audit_pptx(p, floors)
        elif ext == ".docx":
            hits = audit_docx(p, floors)
        elif ext in (".r", ".py"):
            hits = audit_source(p, a.min_src)
        else:
            print(f"-- skip (unsupported): {p.name}")
            continue
        print(f"\n== {p.name}  [{a.profile}] body≥{floors['body']} caption≥{floors['caption']}")
        if not hits:
            print("   OK")
        for where, kind, msg in hits:
            print(f"   {kind:5s} {where:10s} {msg}")
        total += sum(1 for _, k, _ in hits if k != "warn")

    print(f"\n{total} violation(s)  (warn = 캡션이면 의도일 수 있음, 종료코드에 미포함)")
    print("글씨가 안 들어가면 **줄이지 말고 상자·그림을 키운다** (CORE §2①).")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
